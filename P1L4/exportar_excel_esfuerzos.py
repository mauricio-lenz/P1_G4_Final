#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Exporta a Excel los esfuerzos por elemento del edificio completo.

Lee el JSON enriquecido generado por exportar_resultados_unity.py y produce
una hoja de calculo con las fuerzas internas (N, Vy, Vz, T, My, Mz) en los
dos extremos (i, j) de cada elemento y para cada caso/combinacion
(G, Q, EX, EY, C1, C2, C3). Incluye la trazabilidad por tag y la relacion
demanda/capacidad P-M de columnas con curva.

Uso:
  python P1L4/exportar_excel_esfuerzos.py
  python P1L4/exportar_excel_esfuerzos.py --json <ruta> --out <ruta.xlsx>

Requiere: openpyxl (pip install openpyxl).
"""

import argparse
import json
import math
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
JSON_IN = BASE_DIR / "unity_visualizador" / "Assets" / "Resources" / "estructura_p1l4_unity.json"
OUT_DEFAULT = BASE_DIR / "resultados" / "esfuerzos_por_elemento.xlsx"

COMBO_ORDER = ["G", "Q", "EX", "EY", "C1", "C2", "C3"]

FORCE_NAMES = [
    ("N_kN", 0, 6),
    ("Vy_kN", 1, 7),
    ("Vz_kN", 2, 8),
    ("T_kN_m", 3, 9),
    ("My_kN_m", 4, 10),
    ("Mz_kN_m", 5, 11),
]


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def element_meta_map(data):
    meta = {}
    for el in data.get("elements", []):
        meta[int(el["id"])] = el
    return meta


def forces_by_combo(data):
    out = {}
    for combo in COMBO_ORDER:
        out[combo] = {}
    for rec in data.get("p1l4", {}).get("elementForces", []):
        combo = rec.get("combo")
        if combo not in out:
            out[combo] = {}
        out[combo][int(rec["id"])] = [float(x) for x in rec.get("f", [])]
    return out


def pm_capacity_for(data, section_id):
    for curve in data.get("p1l4", {}).get("pmCurves", []):
        if curve.get("sectionId") == section_id:
            return curve
    return None


def moment_capacity_at_p(curve, p_kN):
    """M obtenido de la curva P-M por interpolacion lineal en P."""
    if not curve or not curve.get("points"):
        return None
    points = curve["points"]
    p = p_kN
    for a, b in zip(points, points[1:]):
        pa, pb = float(a["P_kN"]), float(b["P_kN"])
        if min(pa, pb) <= p <= max(pa, pb):
            if abs(pb - pa) < 1e-9:
                return float((a["M_kN_m"] + b["M_kN_m"]) / 2.0)
            t = (p - pa) / (pb - pa)
            return pa * (1 - t) * 0 + float(a["M_kN_m"]) * (1 - t) + float(b["M_kN_m"]) * t
    best = min(points, key=lambda pt: abs(float(pt["P_kN"]) - p))
    return float(best["M_kN_m"])


def main():
    parser = argparse.ArgumentParser(description="Exportar esfuerzos por elemento a Excel")
    parser.add_argument("--json", type=Path, default=JSON_IN, help="JSON de entrada (estructura_p1l4_unity.json)")
    parser.add_argument("--out", type=Path, default=OUT_DEFAULT, help="Archivo .xlsx de salida")
    args = parser.parse_args()

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("ERROR: falta openpyxl. Instala con: pip install openpyxl")
        raise SystemExit(1)

    data = load_json(args.json)
    p1 = data.get("p1l4", {})
    meta = element_meta_map(data)
    forces = forces_by_combo(data)
    combos = [c.get("name") for c in p1.get("combinations", []) if c.get("name")]
    cases = list(COMBO_ORDER)
    combos_byname = {c.get("name"): c for c in p1.get("combinations", []) if c.get("name")}

    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    danger_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    # ── Hoja Elementos (resumen por elemento) ──────────────────────
    ws = wb.active
    ws.title = "Elementos"
    cols = ["ID", "Tag", "Tipo", "Seccion", "Edificio", "Piso",
            "NodoI", "NodoJ", "AreaTrib m2", "CargaTrib kN"]
    ws.append(cols)
    for c in cols:
        ws.cell(row=1, column=cols.index(c) + 1).font = header_font
        ws.cell(row=1, column=cols.index(c) + 1).fill = header_fill
    for el in sorted(data.get("elements", []), key=lambda e: int(e["id"])):
        ws.append([
            int(el["id"]), el.get("elementTag", ""), el.get("type", ""),
            el.get("sectionId") or el.get("seccion", ""), el.get("sourceBuilding", ""),
            el.get("piso", ""), el.get("nodeI", ""), el.get("nodeJ", ""),
            round(float(el.get("areaTributaria", 0.0)), 2),
            round(float(el.get("cargaTributaria", 0.0)), 2),
        ])

    # ── Hojas por caso / combinacion ────────────────────────────────
    for combo in cases:
        label = combos_byname[combo].get("label") if combo in combos_byname else ""
        ws2 = wb.create_sheet(f"Fuerzas {combo}")
        columns = ["ID", "Tag", "Tipo", "Seccion", "Edificio"]
        for comp, i, _ in FORCE_NAMES:
            columns.append(f"{comp} (i)")
        for comp, _, j in FORCE_NAMES:
            columns.append(f"{comp} (j)")
        columns += ["P_compr kN", "M_dem kN*m", "M_cap kN*m", "C = M/Mcap"]
        ws2.append(columns)
        for idx, cname in enumerate(columns, start=1):
            ws2.cell(row=1, column=idx).font = header_font
            ws2.cell(row=1, column=idx).fill = header_fill

        combo_forces = forces.get(combo, {})
        rows = []
        for el in sorted(data.get("elements", []), key=lambda e: int(e["id"])):
            eid = int(el["id"])
            f = combo_forces.get(eid)
            if not f or len(f) < 12:
                continue
            row = [eid, el.get("elementTag", ""), el.get("type", ""),
                   el.get("sectionId") or el.get("seccion", ""), el.get("sourceBuilding", "")]
            p_comp = max(-f[0], -f[6], 0.0)
            m_dem = math.sqrt(f[4] ** 2 + f[5] ** 2)
            m_cap = None
            ratio = ""
            sec_id = el.get("sectionId") or el.get("seccion", "")
            curve = None
            if el.get("type") == "columna":
                curve = pm_capacity_for(data, sec_id) or pm_capacity_for(data, sec_id + "_FIBER")
            if curve:
                m_cap = moment_capacity_at_p(curve, p_comp)
            if m_cap:
                ratio = round(m_dem / m_cap, 3)
            for comp, i, j in FORCE_NAMES:
                row.append(round(f[i], 3))
            for comp, i, j in FORCE_NAMES:
                row.append(round(f[j], 3))
            row += [round(p_comp, 2), round(m_dem, 2),
                    round(m_cap, 2) if m_cap else None, ratio]
            rows.append(row)
        for row in rows:
            ws2.append(row)

        # estilos de encabezados agrupados
        ws2.cell(row=1, column=6).fill = sub_fill
        ws2.cell(row=1, column=12).fill = sub_fill
        for r in range(2, len(rows) + 2):
            ratio_val = ws2.cell(row=r, column=len(columns)).value
            if isinstance(ratio_val, (int, float)) and ratio_val > 1.0:
                for c in range(1, len(columns) + 1):
                    ws2.cell(row=r, column=c).fill = danger_fill
        for idx in range(1, len(columns) + 1):
            letter = get_column_letter(idx)
            ws2.column_dimensions[letter].width = 13 if idx <= 5 else 11

    # ── Hoja Muros ──────────────────────────────────────────────────
    ws3 = wb.create_sheet("Muros")
    wall_cols = ["ID", "Tag", "NodoI", "NodoJ", "t m", "L m", "Bottom", "Top",
                 "Combo", "P kN", "M kN*m", "V kN", "Curva P-M"]
    ws3.append(wall_cols)
    for idx, cname in enumerate(wall_cols, start=1):
        ws3.cell(row=1, column=idx).font = header_font
        ws3.cell(row=1, column=idx).fill = header_fill
    for w in data.get("walls", []):
        for d in w.get("demands", []):
            ws3.append([
                w.get("id", ""), w.get("elementTag", ""), w.get("nodeI", ""),
                w.get("nodeJ", ""), w.get("grosor", ""), w.get("longitud", ""),
                w.get("bottom", ""), w.get("top", ""),
                d.get("combo", ""), d.get("P_kN", ""), d.get("M_kN_m", ""),
                d.get("V_kN", ""), w.get("pmSectionId", ""),
            ])

    # ── Hoja Resumen ────────────────────────────────────────────────
    ws4 = wb.create_sheet("Resumen")
    ws4.append(["Campo", "Valor"])
    ws4["A1"].font = header_font
    ws4["B1"].font = header_font
    ws4["A1"].fill = header_fill
    ws4["B1"].fill = header_fill
    resumen = [
        ("Unidades", data.get("units", "")),
        ("q_G kN/m2", data.get("q_G")),
        ("Q kN/m2", data.get("Q_kN_m2")),
        ("Coef. sismico", data.get("seismic_coefficient")),
        ("Nodos", len(data.get("nodes", []))),
        ("Elementos", len(data.get("elements", []))),
        ("Muros", len(data.get("walls", []))),
        ("Apoyos", len(data.get("supports", []))),
        ("Combinaciones", ", ".join(combos)),
        ("Registros fuerza", sum(len(v) for v in forces.values())),
        ("JSON fuente", str(args.json)),
    ]
    for k, v in resumen:
        ws4.append([k, v])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    counts = ", ".join(f"{c}={len(forces.get(c, {}))}" for c in cases)
    print(f"Excel guardado en: {args.out}")
    print(f"  Elementos: {len(data.get('elements', []))} | registros fuerzas: {counts}")
    print(f"  Muros: {len(data.get('walls', []))}")


if __name__ == "__main__":
    main()