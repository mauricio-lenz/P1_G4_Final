#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""reports/final.md -> reports/final.docx y reports/final.pdf con el formato del informe del curso.

Formato: carta, márgenes de 2,54 cm, Calibri 11 justificado, títulos numerados en mayúscula,
índice general, índice de figuras e índice de tablas con número de página, leyendas
"Tabla N.M:" sobre cada tabla y "Figura N.M:" + "Fuente: Elaboración propia" bajo cada figura.
La página 1 es la portada, con el mismo diseño del informe guía del curso.

El .docx se arma con python-docx; Word (por COM, desde PowerShell) actualiza los índices y
exporta el PDF. Solo funciona en Windows con Microsoft Word instalado.

Uso:
  python -X utf8 Proyecto1/scripts/informe_a_word.py [--md reports/final.md] [--sin-pdf]
"""
import argparse
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
FUENTE = "Calibri"
CODIGO = "Consolas"
AZUL_TABLA = "D9E2F3"
ANCHO_UTIL_CM = 21.59 - 2 * 2.54


# ----------------------------------------------------------------------------
# Estilos
# ----------------------------------------------------------------------------
def configurar_estilos(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2.54))

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for nombre, tam in (("Heading 1", 12), ("Heading 2", 11)):
        st = doc.styles[nombre]
        st.font.name, st.font.size, st.font.bold = FUENTE, Pt(tam), True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for atr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(atr), FUENTE)
        for atr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(atr), None)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for nombre in ("Leyenda Tabla", "Leyenda Figura", "Fuente Figura"):
        st = doc.styles.add_style(nombre, 1)
        st.base_style = normal
        st.font.name, st.font.size = FUENTE, Pt(10 if nombre == "Fuente Figura" else 11)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.space_before = Pt(6 if nombre == "Leyenda Tabla" else 2)
        st.paragraph_format.space_after = Pt(4 if nombre != "Leyenda Figura" else 0)
        st.paragraph_format.keep_with_next = nombre == "Leyenda Tabla" or nombre == "Leyenda Figura"

    cod = doc.styles.add_style("Codigo", 1)
    cod.base_style = normal
    cod.font.name, cod.font.size = CODIGO, Pt(9.5)
    cod.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cod.paragraph_format.space_after = Pt(0)
    cod.paragraph_format.line_spacing = 1.0
    cod.paragraph_format.left_indent = Cm(0.4)

    # entradas del índice: nivel 1 en negrita (como en la guía)
    for nombre, negrita, sangria in (("TOC 1", True, 0.0), ("TOC 2", False, 0.5)):
        try:
            st = doc.styles[nombre]
        except KeyError:
            st = doc.styles.add_style(nombre, 1)
        st.base_style = normal
        st.font.name, st.font.bold = FUENTE, negrita
        st.paragraph_format.left_indent = Cm(sangria)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


# ----------------------------------------------------------------------------
# Utilidades de Word
# ----------------------------------------------------------------------------
def campo(parrafo, instruccion, texto="(actualizar)"):
    """Inserta un campo de Word (TOC, PAGE…) en el párrafo."""
    r = parrafo.add_run()
    ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin"); ini.set(qn("w:dirty"), "true")
    r._r.append(ini)
    r = parrafo.add_run()
    ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = f" {instruccion} "
    r._r.append(ins)
    r = parrafo.add_run()
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r._r.append(sep)
    parrafo.add_run(texto)
    r = parrafo.add_run()
    fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
    r._r.append(fin)


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def repetir_encabezado(fila):
    trpr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trpr.append(el)
    nb = OxmlElement("w:cantSplit"); nb.set(qn("w:val"), "true")
    trpr.append(nb)


def ancho_completo(tabla):
    tblpr = tabla._tbl.tblPr
    w = tblpr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW"); tblpr.append(w)
    w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")


# ----------------------------------------------------------------------------
# Texto en línea: **negrita**, *cursiva*, `código`
# ----------------------------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?!\s)[^*]+?\*(?![\w*]))")


def texto_en_linea(parrafo, texto, negrita=False, cursiva=False, tam=None):
    texto = texto.replace("\\|", "|")
    pos = 0
    for m in TOKEN.finditer(texto):
        if m.start() > pos:
            _run(parrafo, texto[pos:m.start()], negrita, cursiva, tam=tam)
        tok = m.group(0)
        if tok.startswith("**"):
            texto_en_linea(parrafo, tok[2:-2], True, cursiva, tam)
        elif tok.startswith("`"):
            _run(parrafo, tok[1:-1], negrita, cursiva, codigo=True, tam=tam)
        else:
            texto_en_linea(parrafo, tok[1:-1], negrita, True, tam)
        pos = m.end()
    if pos < len(texto):
        _run(parrafo, texto[pos:], negrita, cursiva, tam=tam)


def _run(parrafo, texto, negrita, cursiva, codigo=False, tam=None):
    r = parrafo.add_run(texto)
    r.bold, r.italic = negrita or None, cursiva or None
    if codigo:
        r.font.name = CODIGO
        r.font.size = Pt((tam or 11) - 1.5)
    elif tam:
        r.font.size = Pt(tam)
    return r


# ----------------------------------------------------------------------------
# Bloques
# ----------------------------------------------------------------------------
def agregar_tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    cab, cuerpo = celdas[0], [f for f in celdas[2:]]
    ncol = len(cab)
    tam = 10 if ncol <= 3 else 9.5 if ncol <= 5 else 9
    tabla = doc.add_table(rows=1 + len(cuerpo), cols=ncol)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    ancho_completo(tabla)
    largos = [max(len(re.sub(r"[*`]", "", (f[i] if i < len(f) else ""))) for f in celdas if f is not celdas[1]) for i in range(ncol)]
    total = sum(max(6, l) for l in largos) or 1
    for i, fila in enumerate([cab] + cuerpo):
        for j in range(ncol):
            celda = tabla.cell(i, j)
            celda.width = Cm(ANCHO_UTIL_CM * max(6, largos[j]) / total)
            p = celda.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            texto_en_linea(p, fila[j] if j < len(fila) else "", negrita=(i == 0), tam=tam)
            if i == 0:
                sombrear(celda, AZUL_TABLA)
        if i == 0:
            repetir_encabezado(tabla.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def agregar_figura(doc, ruta, base):
    from PIL import Image
    archivo = (base / ruta).resolve()
    with Image.open(archivo) as im:
        w, h = im.size
    ancho = min(15.5, ANCHO_UTIL_CM)
    alto = ancho * h / w
    if alto > 9.0:
        ancho = 9.0 * w / h
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(archivo), width=Cm(ancho))


def agregar_item(doc, texto, nivel, marca):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.9 + 0.75 * nivel)
    pf.first_line_indent = Cm(-0.5)
    pf.space_after = Pt(3)
    pf.tab_stops.add_tab_stop(Cm(0.9 + 0.75 * nivel))
    _run(p, f"{marca}\t", False, False)
    texto_en_linea(p, texto)


def convertir(md_texto, doc, base):
    cuerpo = md_texto.split("\n---\n", 1)[1] if "\n---\n" in md_texto else md_texto
    lineas = cuerpo.split("\n")
    i, h2, h3 = 0, None, 0
    primer_h1 = True
    referencias = False
    while i < len(lineas):
        ln = lineas[i]
        s = ln.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("## "):
            titulo = s[3:].strip()
            m = re.match(r"(\d+)\.\s*(.*)", titulo)
            h2, h3 = (m.group(1) if m else None), 0
            referencias = not m and titulo.lower().startswith("referencias")
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(0 if primer_h1 else 18)
            primer_h1 = False
            texto_en_linea(p, f"{h2}. {m.group(2).upper()}" if m else titulo.upper())
            i += 1
            continue
        if s.startswith("### "):
            h3 += 1
            p = doc.add_heading(level=2)
            texto_en_linea(p, f"{h2}.{h3} {s[4:].strip()}" if h2 else s[4:].strip())
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            bloque = []
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i]); i += 1
            i += 1
            for k, b in enumerate(bloque):
                p = doc.add_paragraph(style="Codigo")
                p.add_run(b if b else " ")
                if k == 0:
                    p.paragraph_format.space_before = Pt(4)
                if k == len(bloque) - 1:
                    p.paragraph_format.space_after = Pt(6)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2")
                pPr.append(shd)
            continue
        if s.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i]); i += 1
            agregar_tabla(doc, filas)
            continue
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
        if m:
            agregar_figura(doc, m.group(1), base)
            i += 1
            continue
        m = re.match(r"\*(Tabla \d+\.\d+:.*)\*$", s)
        if m:
            p = doc.add_paragraph(style="Leyenda Tabla")
            texto_en_linea(p, m.group(1))
            i += 1
            continue
        m = re.match(r"\*(Figura \d+\.\d+:.*)\*$", s)
        if m:
            p = doc.add_paragraph(style="Leyenda Figura")
            texto_en_linea(p, m.group(1))
            doc.add_paragraph("Fuente: Elaboración propia", style="Fuente Figura")
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        if m:
            nivel = len(m.group(1).replace("\t", "    ")) // 4
            marca = "-" if m.group(2) in "-*" else m.group(2)
            agregar_item(doc, m.group(3), nivel, marca)
            i += 1
            continue
        # párrafo: junta las líneas seguidas
        partes = [s]
        i += 1
        while i < len(lineas) and lineas[i].strip() and not re.match(r"^\s*(#|\||```|!\[|[-*]\s|\d+\.\s|---)", lineas[i]):
            partes.append(lineas[i].strip()); i += 1
        p = doc.add_paragraph()
        if referencias:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(-1.27)
        texto_en_linea(p, " ".join(partes))


PORTADA = {
    "curso": "IOC4102-202620-834 Métodos Computacionales en Obras Civiles",
    "titulo": "Proyecto 1: Laboratorio Estructural OpenSees + Unity + AR",
    "subtitulo": "Segundo Semestre 2026 · Grupo 4",
    "integrantes": ["Matías Campos", "Iván González", "Mauricio Lenz"],
    "profesor": "José Antonio Abell",
    "fecha": "Jueves 08 de Octubre de 2026",
    "logo": "img/final/logo_uandes_fica.png",
}


def _linea(doc, texto="", alin=WD_ALIGN_PARAGRAPH.LEFT, tam=11, negrita=False, despues=0):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = alin
    p.paragraph_format.space_after = Pt(despues)
    if texto:
        r = p.add_run(texto)
        r.font.size, r.bold = Pt(tam), negrita or None
    return p


def portada_e_indices(doc, base, portada):
    """Portada como la del informe guía: curso arriba a la derecha, logo, título centrado,
    integrantes y profesor a la derecha y la fecha abajo, centrada."""
    der, cen = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER
    _linea(doc, portada["curso"], der, 11, despues=18)
    logo = (base / portada["logo"]).resolve()
    if logo.exists():
        _linea(doc).add_run().add_picture(str(logo), width=Cm(11.56))
    for _ in range(7):
        _linea(doc)
    _linea(doc, portada["titulo"], cen, 18, despues=6)
    _linea(doc, portada["subtitulo"], cen, 13)
    for _ in range(9):
        _linea(doc)
    _linea(doc, "Integrantes:", der)
    for nombre in portada["integrantes"]:
        _linea(doc, f"-  {nombre}", der)
    _linea(doc)
    _linea(doc, "Profesor:", der)
    _linea(doc, f"-  {portada['profesor']}", der)
    for _ in range(5):
        _linea(doc)
    p = _linea(doc, portada["fecha"], cen)
    p.add_run().add_break(WD_BREAK.PAGE)

    for titulo, instr, salto in (("ÍNDICE GENERAL", 'TOC \\o "1-2" \\h \\z \\u', True),
                                 ("ÍNDICE DE FIGURAS", 'TOC \\h \\z \\t "Leyenda Figura;1"', False),
                                 ("ÍNDICE DE TABLAS", 'TOC \\h \\z \\t "Leyenda Tabla;1"', False)):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0 if salto else 18)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(titulo); r.bold = True
        campo(doc.add_paragraph(), instr)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def pie_con_numero(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo(p, "PAGE", "1")


def word_a_pdf(docx, pdf):
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{docx}')
  $d.Repaginate()
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Fields.Update() | Out-Null
  $d.Repaginate()
  foreach ($t in $d.TablesOfContents) {{ $t.UpdatePageNumbers() }}
  $d.Save()
  $d.ExportAsFixedFormat('{pdf}', 17)
  $d.Close(0)
}} finally {{ $w.Quit() }}
"""
    subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], check=True, timeout=300)


def main():
    ap = argparse.ArgumentParser(description="Informe final Markdown -> Word y PDF")
    ap.add_argument("--md", type=Path, default=REPO / "reports" / "final.md")
    ap.add_argument("--docx", type=Path, default=None)
    ap.add_argument("--pdf", type=Path, default=None)
    ap.add_argument("--sin-pdf", action="store_true")
    ap.add_argument("--profesor", default=PORTADA["profesor"])
    ap.add_argument("--fecha", default=PORTADA["fecha"])
    a = ap.parse_args()
    md = a.md.resolve()
    docx = (a.docx or md.with_suffix(".docx")).resolve()
    pdf = (a.pdf or md.with_suffix(".pdf")).resolve()

    doc = Document()
    configurar_estilos(doc)
    portada_e_indices(doc, md.parent, {**PORTADA, "profesor": a.profesor, "fecha": a.fecha})
    convertir(md.read_text(encoding="utf-8"), doc, md.parent)
    pie_con_numero(doc)
    doc.save(docx)
    print(docx)
    if not a.sin_pdf:
        word_a_pdf(docx, pdf)
        print(f"{pdf} ({pdf.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
